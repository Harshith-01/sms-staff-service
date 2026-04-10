import logging
import io
import openpyxl
import json

from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile, File
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from sqlalchemy import text
from core.db_errors import db_integrity_http_exception

from core.database import SessionLocal
from core.dependencies import require_role
from core.id_generator import generate_staff_id, generate_user_id
from core.security import get_password_hash
from models.sql_models import (
    NonTeachingStaff, StaffDesignation, User, Role, UserRole,
    Department, TimetableEntry, AcademicTermRef
)
from schemas.dto import (
    StaffCreate, StaffUpdate, StaffOut, StaffListOut,
    StaffDesignationOut, BulkResponse
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/staff", tags=["Staff Service"])

MAX_BULK_FILE_SIZE = 5 * 1024 * 1024
MAX_BULK_ROWS = 500
REQUIRED_BULK_HEADERS = {"full_name", "email", "contact_number", "designation_id", "employment_type", "date_of_joining"}


def _write_audit_log(
    db: Session,
    *,
    action: str,
    entity: str,
    entity_id: str,
    actor_user_id: str,
    old_values: dict | None = None,
    new_values: dict | None = None,
) -> None:
    try:
        db.execute(
            text(
                """
                INSERT INTO public.audit_logs
                (action, entity, entity_id, performed_by, old_values, new_values)
                VALUES (:action, :entity, :entity_id, :performed_by, CAST(:old_values AS JSONB), CAST(:new_values AS JSONB))
                """
            ),
            {
                "action": action,
                "entity": entity,
                "entity_id": entity_id,
                "performed_by": actor_user_id,
                "old_values": None if old_values is None else json.dumps(old_values),
                "new_values": None if new_values is None else json.dumps(new_values),
            },
        )
    except Exception:
        logger.exception("staff_audit_log_failed action=%s entity=%s entity_id=%s", action, entity, entity_id)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# ============================================================
# INTERNAL LOGIC
# ============================================================

def create_staff_logic(db: Session, data: StaffCreate) -> str:
    normalized_email = data.email.strip().lower()

    if db.query(User).filter(User.email == normalized_email).first():
        raise HTTPException(409, "Email already registered")

    if db.query(NonTeachingStaff).filter(NonTeachingStaff.email == normalized_email).first():
        raise HTTPException(409, "Staff email already exists")

    if data.designation_id and not db.get(StaffDesignation, data.designation_id):
        raise HTTPException(404, "Designation not found")

    if data.department_id and not db.get(Department, data.department_id):
        raise HTTPException(404, "Department not found")

    staff_id = generate_staff_id(db)
    user_id = generate_user_id(db)

    temp_password = data.contact_number or "Welcome@123"

    new_user = User(
        id=user_id,
        email=normalized_email,
        hashed_password=get_password_hash(temp_password),
        is_active=True,
    )
    db.add(new_user)
    db.flush()

    role = db.query(Role).filter(Role.role_name == "NON_TEACHING_STAFF").first()
    if not role:
        raise HTTPException(500, "NON_TEACHING_STAFF role not found in roles table")

    db.add(UserRole(user_id=user_id, role_id=role.id))

    staff = NonTeachingStaff(
        id=staff_id,
        user_id=user_id,
        full_name=data.full_name.strip(),
        designation_id=data.designation_id,
        department_id=data.department_id,
        email=normalized_email,
        contact_number=data.contact_number,
        emergency_contact=data.emergency_contact,
        address=data.address,
        date_of_joining=data.date_of_joining,
        employment_type=data.employment_type,
        status="ACTIVE",
        date_of_birth=data.date_of_birth,
        gender=data.gender,
        aadhaar_number=data.aadhaar_number,
    )
    db.add(staff)
    return staff_id


# ============================================================
# CREATE
# ============================================================

@router.post("", status_code=201)
def create_staff(
    data: StaffCreate,
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    try:
        staff_id = create_staff_logic(db, data)
        _write_audit_log(
            db,
            action="CREATE",
            entity="NonTeachingStaff",
            entity_id=staff_id,
            actor_user_id=user.get("user_id", "unknown"),
            new_values={"email": data.email, "designation_id": data.designation_id, "department_id": data.department_id},
        )
        db.commit()
        logger.info(f"Staff created: {staff_id} by {user.get('user_id')}")
        return {"staff_id": staff_id}
    except IntegrityError as exc:
        db.rollback()
        raise db_integrity_http_exception(exc, fallback_status=409, fallback_detail="Duplicate or invalid staff data")
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        logger.exception("Staff creation failed")
        raise HTTPException(500, "Internal server error")


# ============================================================
# BULK CREATE
# ============================================================

@router.post("/bulk", response_model=BulkResponse)
async def bulk_create_staff(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    contents = await file.read()
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Only .xlsx files allowed")
    if len(contents) > MAX_BULK_FILE_SIZE:
        raise HTTPException(400, "File too large (max 5MB)")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Invalid or corrupted Excel file")

    sheet = wb.active
    if sheet.max_row > MAX_BULK_ROWS:
        raise HTTPException(400, "Maximum 500 rows allowed")

    headers = [cell.value for cell in sheet[1]]
    header_map = {str(h).strip(): i for i, h in enumerate(headers) if h}
    missing = REQUIRED_BULK_HEADERS - set(header_map.keys())
    if missing:
        raise HTTPException(400, f"Missing headers: {', '.join(sorted(missing))}")

    def get_val(row, name):
        try:
            return row[header_map[name]].value
        except Exception:
            return None

    success, errors = 0, []

    for i, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        try:
            with db.begin_nested():
                from schemas.dto import StaffCreate as SC
                from datetime import date as _date
                doj_raw = get_val(row, "date_of_joining")
                doj = doj_raw if isinstance(doj_raw, _date) else None

                data = SC(
                    full_name=get_val(row, "full_name"),
                    email=get_val(row, "email"),
                    contact_number=str(get_val(row, "contact_number") or "").strip() or None,
                    designation_id=int(get_val(row, "designation_id")) if get_val(row, "designation_id") else None,
                    department_id=int(get_val(row, "department_id")) if get_val(row, "department_id") else None,
                    employment_type=get_val(row, "employment_type") or "FULL_TIME",
                    date_of_joining=doj,
                )
                create_staff_logic(db, data)
                success += 1
        except Exception as e:
            errors.append(f"Row {i}: {str(e)}")

    db.commit()
    return BulkResponse(success_count=success, error_count=len(errors), errors=errors)


# ============================================================
# LIST
# ============================================================

@router.get("", response_model=list[StaffListOut])
def list_staff(
    designation_id: int | None = None,
    department_id: int | None = None,
    status: str | None = None,
    full_name: str | None = None,
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    query = db.query(NonTeachingStaff)

    if status:
        query = query.filter(NonTeachingStaff.status == status)
    else:
        query = query.filter(NonTeachingStaff.status == "ACTIVE")

    if designation_id:
        query = query.filter(NonTeachingStaff.designation_id == designation_id)
    if department_id:
        query = query.filter(NonTeachingStaff.department_id == department_id)
    if full_name:
        query = query.filter(NonTeachingStaff.full_name.ilike(f"%{full_name}%"))

    results = query.offset(offset).limit(limit).all()
    return [StaffListOut.model_validate(s) for s in results]


# ============================================================
# SELF PROFILE
# ============================================================

@router.get("/me", response_model=StaffOut)
def get_my_profile(
    db: Session = Depends(get_db),
    user=Depends(require_role(["NON_TEACHING_STAFF"])),
):
    staff = db.query(NonTeachingStaff).filter(
        NonTeachingStaff.user_id == user["user_id"]
    ).first()
    if not staff:
        raise HTTPException(404, "Staff profile not found")
    return StaffOut.model_validate(staff)


# ============================================================
# TIMETABLE VIEW (staff can read school timetable)
# ============================================================

@router.get("/me/timetable")
def view_timetable(
    class_section_id: int,
    academic_term_id: int,
    db: Session = Depends(get_db),
    user=Depends(require_role(["NON_TEACHING_STAFF", "ADMIN"])),
):
    entries = (
        db.query(TimetableEntry)
        .filter(
            TimetableEntry.class_section_id == class_section_id,
            TimetableEntry.academic_term_id == academic_term_id,
            TimetableEntry.is_active.is_(True),
        )
        .order_by(TimetableEntry.weekday.asc(), TimetableEntry.period_no.asc())
        .all()
    )
    return [
        {
            "id": e.id,
            "weekday": e.weekday,
            "period_no": e.period_no,
            "subject_id": e.subject_id,
            "teacher_id": e.teacher_id,
            "room_code": e.room_code,
            "entry_type": e.entry_type,
        }
        for e in entries
    ]


# ============================================================
# GET BY ID (admin only)
# ============================================================

@router.get("/{staff_id}", response_model=StaffOut)
def get_staff(
    staff_id: str,
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    staff = db.query(NonTeachingStaff).filter(NonTeachingStaff.id == staff_id).first()
    if not staff:
        raise HTTPException(404, "Staff member not found")
    return StaffOut.model_validate(staff)


# ============================================================
# UPDATE
# ============================================================

@router.put("/{staff_id}")
def update_staff(
    staff_id: str,
    data: StaffUpdate,
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    staff = db.query(NonTeachingStaff).filter(NonTeachingStaff.id == staff_id).first()
    if not staff:
        raise HTTPException(404, "Staff member not found")

    updates = data.model_dump(exclude_unset=True)

    if "email" in updates:
        new_email = updates["email"].strip().lower()
        existing = db.query(User).filter(
            User.email == new_email,
            User.id != staff.user_id,
        ).first()
        if existing:
            raise HTTPException(409, "Email already in use")
        auth_user = db.query(User).filter(User.id == staff.user_id).first()
        if auth_user:
            auth_user.email = new_email
        staff.email = new_email
        updates.pop("email", None)

    if "status" in updates:
        auth_user = db.query(User).filter(User.id == staff.user_id).first()
        if auth_user:
            auth_user.is_active = updates["status"] == "ACTIVE"

    for key, value in updates.items():
        setattr(staff, key, value)

    try:
        _write_audit_log(
            db,
            action="UPDATE",
            entity="NonTeachingStaff",
            entity_id=staff_id,
            actor_user_id=user.get("user_id", "unknown"),
            new_values=updates,
        )
        db.commit()
        return {"message": "Staff updated successfully"}
    except IntegrityError as exc:
        db.rollback()
        raise db_integrity_http_exception(exc, fallback_status=409, fallback_detail="Update constraint violation")


# ============================================================
# SOFT DELETE / DEACTIVATE
# ============================================================

@router.delete("/{staff_id}")
def deactivate_staff(
    staff_id: str,
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN"])),
):
    staff = db.query(NonTeachingStaff).filter(NonTeachingStaff.id == staff_id).first()
    if not staff:
        raise HTTPException(404, "Staff member not found")

    from datetime import datetime, timezone
    staff.status = "INACTIVE"
    staff.deleted_at = datetime.now(timezone.utc)

    auth_user = db.query(User).filter(User.id == staff.user_id).first()
    if auth_user:
        auth_user.is_active = False

    _write_audit_log(
        db,
        action="DELETE",
        entity="NonTeachingStaff",
        entity_id=staff_id,
        actor_user_id=user.get("user_id", "unknown"),
        old_values={"status": "ACTIVE"},
        new_values={"status": "INACTIVE"},
    )
    db.commit()
    return {"message": "Staff deactivated"}


# ============================================================
# DESIGNATIONS (reference lookup)
# ============================================================

@router.get("/designations/list", response_model=list[StaffDesignationOut])
def list_designations(
    db: Session = Depends(get_db),
    user=Depends(require_role(["ADMIN", "NON_TEACHING_STAFF"])),
):
    return db.query(StaffDesignation).filter(StaffDesignation.is_active.is_(True)).all()
